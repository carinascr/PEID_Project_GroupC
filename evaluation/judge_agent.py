"""Golden-set execution and one-criterion LLM judging."""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from src.assistant import answer_question
from src.contracts import Contract
from src.llm import ask


ROOT = Path(__file__).parents[1]
GOLDEN_PATH = ROOT / "evaluation" / "goldenset.xlsx"
JUDGE_PROMPT_PATH = ROOT / "evaluation" / "judge_prompt.md"


@dataclass(frozen=True)
class GoldenCase:
    case_id: str
    question: str
    expected_answer: str


@dataclass
class CaseEvidence:
    case_id: str
    question: str
    expected_answer: str
    candidate_answer: str
    candidate_record_ids: list[str]
    deterministic_status: str
    judge_verdict: str
    judge_evidence: str
    judge_reason: str


def load_golden_cases(path: Path = GOLDEN_PATH) -> list[GoldenCase]:
    """Load only populated rows and preserve workbook values as text."""
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    cases = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        values = ["" if value is None else str(value).strip() for value in row[:3]]
        if not any(values):
            continue
        if len(values) != 3 or not all(values):
            raise ValueError(f"Incomplete golden-set row: {values!r}")
        cases.append(GoldenCase(values[0], values[1], values[2]))
    if not cases:
        raise ValueError(f"No populated cases found in {path}")
    return cases


def _normalise(value: str) -> str:
    return re.sub(r"\s+", " ", value.casefold().strip())


def deterministic_status(expected: str, candidate: str) -> str:
    """Check exact supported answers without asking an LLM."""
    if expected == "#VALUE!":
        return "unknown"
    expected_normalised = _normalise(expected)
    candidate_normalised = _normalise(candidate)
    if expected_normalised in candidate_normalised:
        return "pass"

    def numeric_values(value: str) -> list[str]:
        values = []
        for token in re.findall(r"(?<!\w)(?:€\s*)?[\d][\d.,]*(?!\w)", value):
            number = token.replace("€", "").replace(" ", "")
            if "," in number and "." in number:
                number = number.replace(",", "") if number.rfind(".") > number.rfind(",") else number.replace(".", "").replace(",", ".")
            elif "," in number:
                number = number.replace(",", ".")
            try:
                values.append(f"{float(number):.2f}")
            except ValueError:
                continue
        return values

    expected_numbers = numeric_values(expected)
    candidate_numbers = numeric_values(candidate)
    if expected_numbers and all(number in candidate_numbers for number in expected_numbers):
        return "pass"
    return "fail"


def _parse_judge_response(response: str) -> dict[str, str]:
    try:
        parsed = json.loads(response)
    except json.JSONDecodeError as exc:
        raise ValueError("Judge returned invalid JSON") from exc
    required = {"verdict", "evidence", "reason"}
    if not isinstance(parsed, dict) or not required <= parsed.keys():
        raise ValueError("Judge response is missing required fields")
    verdict = str(parsed["verdict"]).upper()
    if verdict not in {"PASS", "FAIL", "UNKNOWN"}:
        raise ValueError(f"Judge returned invalid verdict: {verdict}")
    return {
        "verdict": verdict,
        "evidence": str(parsed["evidence"]),
        "reason": str(parsed["reason"]),
    }


def judge_case(
    case: GoldenCase,
    candidate: dict[str, Any],
    judge: Callable[..., str] = ask,
) -> dict[str, str]:
    """Judge one criterion for one case through the shared LLM client."""
    prompt = (
        f"Question:\n{case.question}\n\n"
        f"Golden answer:\n{case.expected_answer}\n\n"
        f"Candidate answer:\n{candidate.get('answer', 'unknown')}"
    )
    return _parse_judge_response(
        judge(prompt=prompt, system=JUDGE_PROMPT_PATH.read_text(encoding="utf-8"))
    )


def run_evaluation(
    contracts: list[Contract],
    *,
    use_llm_judge: bool = False,
    judge: Callable[..., str] = ask,
    limit: int | None = None,
) -> list[CaseEvidence]:
    evidence = []
    cases = load_golden_cases()
    if limit is not None:
        if limit < 1:
            raise ValueError("Evaluation limit must be greater than zero.")
        cases = cases[:limit]
    for case in cases:
        candidate = answer_question(contracts, case.question)
        candidate_answer = str(candidate.get("answer", "unknown"))
        if case.expected_answer == "#VALUE!":
            judged = {
                "verdict": "UNKNOWN",
                "evidence": "unknown",
                "reason": "Golden set marks this case as #VALUE!, so it is not judge-scorable.",
            }
        elif use_llm_judge:
            try:
                judged = judge_case(case, candidate, judge)
            except Exception as exc:
                judged = {
                    "verdict": "UNKNOWN",
                    "evidence": "unknown",
                    "reason": f"Judge request failed: {type(exc).__name__}: {exc}",
                }
        else:
            judged = {
                "verdict": deterministic_status(case.expected_answer, candidate_answer).upper(),
                "evidence": candidate_answer[:240] or "unknown",
                "reason": "Deterministic golden-answer comparison.",
            }
        evidence.append(
            CaseEvidence(
                case_id=case.case_id,
                question=case.question,
                expected_answer=case.expected_answer,
                candidate_answer=candidate_answer,
                candidate_record_ids=list(candidate.get("record_ids", [])),
                deterministic_status=deterministic_status(
                    case.expected_answer, candidate_answer
                ),
                judge_verdict=judged["verdict"],
                judge_evidence=judged["evidence"],
                judge_reason=judged["reason"],
            )
        )
    return evidence


def calculate_metrics(evidence: list[CaseEvidence]) -> dict[str, float | int]:
    """Calculate answer-level accuracy, precision, and recall.

    UNKNOWN cases are excluded from the binary denominator because the
    workbook explicitly marks them as not scorable.
    """
    scorable = [item for item in evidence if item.deterministic_status != "unknown"]
    judged = [item for item in scorable if item.judge_verdict in {"PASS", "FAIL"}]
    actual_positive = sum(item.deterministic_status == "pass" for item in scorable)
    predicted_positive = sum(item.judge_verdict == "PASS" for item in judged)
    true_positive = sum(
        item.deterministic_status == "pass" and item.judge_verdict == "PASS"
        for item in judged
    )
    true_negative = sum(
        item.deterministic_status == "fail" and item.judge_verdict == "FAIL"
        for item in judged
    )
    predicted_negative = sum(item.judge_verdict == "FAIL" for item in judged)
    accuracy = (true_positive + true_negative) / len(judged) if judged else 0.0
    precision = true_positive / predicted_positive if predicted_positive else 0.0
    recall = true_positive / actual_positive if actual_positive else 0.0
    return {
        "cases": len(evidence),
        "scorable_cases": len(scorable),
        "unknown_cases": len(evidence) - len(scorable),
        "judge_unknown_cases": len(scorable) - len(judged),
        "answer_pass_rate": actual_positive / len(scorable) if scorable else 0.0,
        "true_positive": true_positive,
        "true_negative": true_negative,
        "false_positive": predicted_positive - true_positive,
        "false_negative": predicted_negative - true_negative,
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
    }


def evidence_to_json(evidence: list[CaseEvidence]) -> list[dict[str, Any]]:
    return [asdict(item) for item in evidence]
