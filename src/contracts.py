"""Load and search the contract source while preserving provenance."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
import unicodedata

try:
    from .vector_store import ContractVectorStore
except ImportError:
    from vector_store import ContractVectorStore


SOURCE_COLUMNS = (
    "Objeto do Contrato",
    "Tipo de Procedimento",
    "Tipo(s) de Contrato",
    "CPV",
    "CPV Tipo",
    "CPV Designação",
    "Entidade(s) Adjudicante(s)",
    "Entidade(s) Adjudicatária(s)",
    "Preço Contratual",
    "Data de Publicação",
    "Data de Celebração do Contrato",
    "Prazo de Execução",
    "Local de Execução",
    "Fundamentação",
    "Estado",
    "Regime",
    "Critérios Materiais",
    "Procedimento Centralizado",
)


@dataclass(frozen=True)
class Contract:
    record_id: str
    source_row: int
    fields: dict[str, str]
    text: str


def _clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "nat", "none"} else text


def _key(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^\w]+", " ", text).strip()


def _rows_from_xlsx(path: Path) -> Iterable[dict[str, object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    headers = [_clean(value) for value in next(values)]
    for row in values:
        yield dict(zip(headers, row))


def _rows_from_csv(path: Path) -> Iterable[dict[str, object]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Unable to decode dataset: {path}")
    yield from csv.DictReader(io.StringIO(text))


def _source_rows(path: Path) -> Iterable[dict[str, object]]:
    if zipfile.is_zipfile(path):
        yield from _rows_from_xlsx(path)
    else:
        yield from _rows_from_csv(path)


def load_contracts(path: str | Path) -> list[Contract]:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"Contract dataset not found: {source}")
    contracts = []
    for row_number, row in enumerate(_source_rows(source), start=2):
        fields = {key: _clean(row.get(key, "")) for key in SOURCE_COLUMNS}
        text = "\n".join(
            f"{key}: {value}" for key, value in fields.items() if value
        )
        contracts.append(
            Contract(
                record_id=f"contrato-{row_number - 1:04d}",
                source_row=row_number,
                fields=fields,
                text=text,
            )
        )
    if not contracts:
        raise ValueError(f"Contract dataset is empty: {source}")
    return contracts


def search_contracts(
    contracts: list[Contract], query: str, limit: int = 5
) -> list[tuple[Contract, float]]:
    if not query.strip():
        return []
    query_terms = set(_key(query).split())
    normalized_query = _key(query)
    ranking_query = (
        any(term in normalized_query for term in ("maior", "maiores"))
        and ("contrat" in normalized_query or "preco" in normalized_query)
    )
    if ranking_query:
        requested_limit = re.search(r"\b(\d+)\b", normalized_query)
        result_limit = int(requested_limit.group(1)) if requested_limit else limit
        ranked = []
        for contract in contracts:
            raw_price = contract.fields.get("Preço Contratual", "").replace(",", ".")
            try:
                price = float(re.sub(r"[^\d.-]", "", raw_price))
            except ValueError:
                continue
            ranked.append((contract, price))
        return sorted(ranked, key=lambda item: (-item[1], item[0].record_id))[:result_limit]
    vector_store = ContractVectorStore(
        (contract.record_id, contract.text) for contract in contracts
    )
    vector_scores = dict(vector_store.search(query, limit=len(contracts)))
    scored = []
    for contract in contracts:
        terms = set(_key(contract.text).split())
        overlap = len(query_terms & terms)
        phrase_bonus = 2.0 if _key(query) in _key(contract.text) else 0.0
        lexical_score = overlap / max(len(query_terms), 1) + phrase_bonus
        vector_score = vector_scores.get(contract.record_id, 0.0)
        score = lexical_score + vector_score
        if score:
            scored.append((contract, score))
    return sorted(scored, key=lambda item: (-item[1], item[0].record_id))[:limit]


def format_context(results: list[tuple[Contract, float]]) -> str:
    return "\n\n".join(
        f"[{contract.record_id} | source row {contract.source_row}]\n{contract.text}"
        for contract, _ in results
    )
